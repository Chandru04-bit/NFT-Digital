import os
import re
import openpyxl

html_files = [f for f in os.listdir('.') if f.endswith('.html')]
print(f"Checking {len(html_files)} HTML files: {html_files}")

all_valid = True
for html in html_files:
    with open(html, 'r', encoding='utf-8') as f:
        content = f.read()
    if '</html>' not in content:
        print(f"Error: Missing </html> in {html}")
        all_valid = False
    
    # Check src / href links
    local_links = re.findall(r'(?:src|href)=["\']([^"\']+)["\']', content)
    for link in local_links:
        if not link.startswith('http') and not link.startswith('#') and not link.startswith('mailto:') and link != '':
            path = link.split('?')[0]
            if not os.path.exists(path):
                print(f"Warning: Link not found in {html}: {path}")
                all_valid = False

if all_valid:
    print("All local assets and links in HTML files exist!")

wb = openpyxl.load_workbook("NFT_Marketplace_Daily_Progress.xlsx")
ws = wb.active
print(f"Excel validated: Sheet title '{ws.title}', {ws.max_row} rows, {ws.max_column} columns.")
