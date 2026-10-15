import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Daily Progress Report"

# Ensure grid lines are visible
ws.views.sheetView[0].showGridLines = True

# Title block
ws.merge_cells("A1:J1")
ws["A1"] = "NFT DIGITAL MARKETPLACE - DAILY PROGRESS TRACKING REPORT"
title_font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
title_fill = PatternFill(start_color="121025", end_color="121025", fill_type="solid")
ws["A1"].font = title_font
ws["A1"].fill = title_fill
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 40

# Subtitle block
ws.merge_cells("A2:J2")
ws["A2"] = "Project: NFT Digital Marketplace Dashboard | Tech: HTML5, CSS3, JS (ES6+), Bootstrap 5, Chart.js, LocalStorage"
ws["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="D1D5DB")
ws["A2"].fill = PatternFill(start_color="1F1B3E", end_color="1F1B3E", fill_type="solid")
ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 24

# Header row
headers = [
    "Date", "Day", "Task / Module", "Planned Work", "Completed Work",
    "Pending Work", "Issues / Bugs", "Status", "Progress %", "Remarks"
]
ws.row_dimensions[4].height = 28

header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="735BF2", end_color="735BF2", fill_type="solid")
thin_border = Border(
    left=Side(style='thin', color='E5E7EB'),
    right=Side(style='thin', color='E5E7EB'),
    top=Side(style='thin', color='E5E7EB'),
    bottom=Side(style='thin', color='E5E7EB')
)

for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

# Data Rows for Day 1, Day 2, Day 3
data = [
    # Day 1
    ("2026-10-14", "Day 1", "UI Architecture & Design System",
     "Analyze Figma screens, define CSS variables (#121025 background, #1e1b38 cards, #735bf2 accents), set up typography & shared assets",
     "Complete design tokens, folder structure, 12 vector NFT artworks, creator avatars, responsive grid & CSS components",
     "None", "None", "Completed", "100%", "Dark theme foundation established with smooth light theme compatibility"),
    
    ("2026-10-14", "Day 1", "Shared Navigation & Header",
     "Build sticky desktop sidebar, mobile off-canvas drawer, top search bar with live dropdown, theme toggle & profile menu",
     "Sidebar with active links, responsive header, notification bell with badge & dropdown, instant NFT search, theme switcher",
     "None", "None", "Completed", "100%", "Unified navigation across all pages; smooth collapse on mobile viewports"),
    
    ("2026-10-14", "Day 1", "Home Dashboard (index.html)",
     "Implement hero banner, trending bids filter (All/Artwork/Book), Chart.js ETH price & Doughnut stats, top creators, active bids table",
     "Full hero with live featured countdown, category filter pills, interactive Chart.js charts, follow/unfollow creators, active bids",
     "None", "None", "Completed", "100%", "Charts styled with purple gradient; live countdown updates every second"),

    ("2026-10-14", "Day 1", "Bids Dashboard (bids.html)",
     "Create bids summary KPI cards (Artworks, Auctions, Creators, Cancelled), active bids data table with checkboxes and actions",
     "4 KPI stat cards, full interactive bids table with countdowns, row selection, row cancel action, and Place Bid trigger",
     "None", "None", "Completed", "100%", "Connects directly to universal bid modal with amount validation"),

    ("2026-10-14", "Day 1", "Saved Items (saved.html)",
     "Build wishlist grid, filter by category/sorting, remove from saved, empty state with illustration & explore CTA",
     "Saved items grid loaded from LocalStorage, toggle save heart icon, instant removal with toast, responsive empty state",
     "None", "None", "Completed", "100%", "State synchronizes seamlessly with Home and Collections"),

    ("2026-10-14", "Day 1", "Collections (collections.html)",
     "Build collections page with category filter (All, Artwork, Book, 3D Art, Metaverse), search input, sorting dropdown & details modal",
     "Responsive NFT card grid, instant search filter, price/recency sorting, category pill buttons, bid modal launch",
     "None", "None", "Completed", "100%", "Consistent card heights, image ratios, and badge styling"),

    # Day 2
    ("2026-10-15", "Day 2", "Profile Management (profile.html)",
     "User header with banner, verification badges, stats counters, tabs for My Bought, My Collections, Following list, Edit Profile modal",
     "Complete profile layout, verified creator badge, 2FA status, dynamic bought/collected tabs, follow/unfollow toggle, edit profile persistence",
     "None", "None", "Completed", "100%", "LocalStorage stores edited full name, username, bio, and avatar instantly"),

    ("2026-10-15", "Day 2", "Settings Multi-Tab (settings.html)",
     "Create 6 settings tabs: Profile, Application, Security, Activity, Payment Method, API key management without full page reloads",
     "All 6 tabs functional: theme/notification preferences, password form with strength check, session table, demo card/wallet, API key mask/copy",
     "None", "None", "Completed", "100%", "Clear simulated demo badges for wallets and API security"),

    ("2026-10-15", "Day 2", "Authentication (login & register.html)",
     "Create sleek login and register pages with form validation, show/hide password, demo credentials button, session persistence",
     "Complete dark theme auth screens, email/password validation, session state stored in LocalStorage, logout flow to login.html",
     "None", "None", "Completed", "100%", "Quick-fill demo account button provided for immediate testing"),

    ("2026-10-15", "Day 2", "NFT Details Page & Universal Modal",
     "Standalone nft-details.html and reusable modal with high-res art, attributes, creator info, countdown, bid history, bidding form",
     "Full NFT details page with dynamic query param loading, tabbed history & attributes, plus modal triggerable from every card",
     "None", "None", "Completed", "100%", "Validates bid amounts, verifies minimum bid rule, logs transaction history"),

    ("2026-10-15", "Day 2", "Data Store & LocalStorage Architecture",
     "Centralized data engine in js/data.js managing NFTs, bids, saved items, profile data, activity logs, and theme preferences",
     "Reactive data store with automated fallback seeding, bid placement logic, wishlist toggling, and multi-tab sync",
     "None", "None", "Completed", "100%", "Zero code duplication across pages"),

    # Day 3
    ("2026-10-16", "Day 3", "Responsive Optimization (320px - 1920px)",
     "Test and polish across 16 specified viewports: 320, 360, 375, 390, 414, 480, 576, 768, 820, 1024, 1152, 1280, 1366, 1440, 1600, 1920px",
     "Off-canvas mobile menu, horizontal scrolling tables with touch support, auto-fitting CSS grid, flexible hero typography, zero layout overflow",
     "None", "None", "Completed", "100%", "Passed all breakpoint checks cleanly"),

    ("2026-10-16", "Day 3", "Quality Assurance & Interactive Polish",
     "Validate all forms, keyboard focus outlines, toast feedback, countdown timers, search highlights, dark/light theme persistence",
     "All interactive triggers tested, console error free, toast notifications for all user actions, comprehensive README documentation",
     "None", "None", "Completed", "100%", "Final checklist 100% verified against project specifications")
]

row_start = 5
for row_idx, item in enumerate(data, start=row_start):
    ws.row_dimensions[row_idx].height = 24
    is_even = (row_idx % 2 == 0)
    bg_color = "F9FAFB" if is_even else "FFFFFF"
    row_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    
    for col_idx, val in enumerate(item, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = Font(name="Segoe UI", size=9, color="1F2937")
        cell.fill = row_fill
        cell.border = thin_border
        
        # Center align dates, days, status, progress
        if col_idx in [1, 2, 8, 9]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            
        # Highlight Status
        if col_idx == 8:
            cell.font = Font(name="Segoe UI", size=9, bold=True, color="047857")
            cell.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
        elif col_idx == 9:
            cell.font = Font(name="Segoe UI", size=9, bold=True, color="4338CA")
            cell.fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")

# Adjust column widths
col_widths = {
    1: 14,  # Date
    2: 10,  # Day
    3: 32,  # Task/Module
    4: 42,  # Planned Work
    5: 42,  # Completed Work
    6: 16,  # Pending Work
    7: 16,  # Issues/Bugs
    8: 14,  # Status
    9: 14,  # Progress %
    10: 42  # Remarks
}

for col_idx, width in col_widths.items():
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = width

wb.save("NFT_Marketplace_Daily_Progress.xlsx")
print("Excel sheet generated successfully at NFT_Marketplace_Daily_Progress.xlsx")
